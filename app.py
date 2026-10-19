import os
import io
import csv
import uuid
import hashlib
import datetime
import difflib
import webbrowser
from threading import Timer

from flask import Flask, render_template, request, Response
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.protection import SheetProtection
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

# ============================================================
# ENTERPRISE HASH AUDIT STUDIO — PUBLIC DEMO
# Demo-only copy. Maximum 50 PDFs per operation.
# Do not add permanent-license checks to this demo build.
# ============================================================

DEMO_LABEL = "DEMO SOFTWARE"
MAX_PDFS = 50
MAX_CONTENT_LENGTH = 1000 * 1024 * 1024
PORT = int(os.environ.get("PORT", "28888"))
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
QUARANTINE_DIR = os.path.join(BASE_DIR, "Quarantine_Folder")

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("DEMO_SECRET_KEY", "demo-only-local-key-change-before-hosting")
app.config["MAX_CONTENT_LENGTH"] = MAX_CONTENT_LENGTH

LATEST_PROCESSED_FILES = []
MEMORI_TEKS_ASLI = {}


def hitung_sha256_file(file_storage):
    """Return the SHA-256 digest of a Flask uploaded file without saving it."""
    digest = hashlib.sha256()
    file_storage.seek(0)
    for block in iter(lambda: file_storage.read(1024 * 1024), b""):
        digest.update(block)
    file_storage.seek(0)
    return digest.hexdigest()


def ekstrak_teks_pdf(file_storage):
    teks = ""
    try:
        file_storage.seek(0)
        reader = PdfReader(file_storage)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                teks += page_text + "\n"
    except Exception as exc:
        teks = f"[Gagal ekstrak teks: {exc}]"
    finally:
        file_storage.seek(0)
    return teks


def proses_guid_dan_ekstrak_teks(file_storage, idx):
    nama_asli = secure_filename(file_storage.filename or f"document_{idx}.pdf")
    kode_guid = str(uuid.uuid4())
    MEMORI_TEKS_ASLI[kode_guid] = ekstrak_teks_pdf(file_storage)
    return {
        "id": idx,
        "nama_asli": nama_asli,
        "nama_guid": f"{kode_guid}.pdf",
        "guid_key": kode_guid,
    }


def batasi_pdf(files):
    valid_files = [
        item for item in files
        if item and item.filename and item.filename.lower().endswith(".pdf")
    ]
    if len(valid_files) > MAX_PDFS:
        return None, f"Demo dibatasi maksimal {MAX_PDFS} PDF per proses. File yang dipilih: {len(valid_files)}."
    return valid_files, None


def buat_watermark_png(path):
    """Create a repeated diagonal watermark image for embedding in Excel."""
    width, height = 1500, 1000
    base = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    layer = Image.new("RGBA", (900, 150), (255, 255, 255, 0))
    draw = ImageDraw.Draw(layer)
    font = None
    for candidate in ["arial.ttf", "Arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"]:
        try:
            font = ImageFont.truetype(candidate, 48)
            break
        except Exception:
            pass
    if font is None:
        font = ImageFont.load_default()
    draw.text((20, 35), DEMO_LABEL, fill=(190, 0, 0, 80), font=font, stroke_width=1)
    rotated = layer.rotate(35, expand=True, resample=Image.Resampling.BICUBIC)
    for y in range(-250, height + 200, 230):
        for x in range(-500, width, 650):
            base.paste(rotated, (x, y), rotated)
    base.save(path, format="PNG")


def pasang_watermark_excel(workbook, sheet_names=None, protect=True):
    """Embed repeated diagonal watermark in each target worksheet."""
    watermark_path = os.path.join(BASE_DIR, "_demo_watermark.png")
    buat_watermark_png(watermark_path)
    targets = sheet_names or workbook.sheetnames
    for sheet_name in targets:
        ws = workbook[sheet_name]
        # A visible banner remains searchable even when images are not displayed.
        ws.insert_rows(1, 1)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(4, ws.max_column))
        banner = ws.cell(row=1, column=1, value=f"{DEMO_LABEL} — FOR DEMONSTRATION ONLY — MAXIMUM {MAX_PDFS} PDF")
        banner.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        banner.fill = PatternFill(fill_type="solid", fgColor="B91C1C")
        banner.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        img = XLImage(watermark_path)
        img.width = 750
        img.height = 500
        ws.add_image(img, "A2")
        ws.sheet_view.showGridLines = True
        if protect:
            ws.protection.sheet = True
            ws.protection.objects = True
            ws.protection.set_password("DEMO-SOFTWARE-50PDF")
            ws.protection.selectLockedCells = True
            ws.protection.selectUnlockedCells = True
    return watermark_path


def jalankan_audit_forensik_massal(daftar_file_diuji, hash_master_set):
    os.makedirs(QUARANTINE_DIR, exist_ok=True)
    data_detail_laporan = []
    baris_dikarantina = []
    total_aman = 0

    for idx, uploaded in enumerate(daftar_file_diuji, start=1):
        nama_asli = secure_filename(uploaded.filename or f"document_{idx}.pdf")
        uploaded.seek(0)
        hash_sekarang = hitung_sha256_file(uploaded).lower().strip()
        teks_baru = ekstrak_teks_pdf(uploaded)

        if hash_sekarang in hash_master_set:
            status = "✓ VALID"
            lokasi_sekarang = "Uploaded in demo session (not saved)"
            detail_forensik = "No SHA-256 mismatch against the supplied master list."
            total_aman += 1
        else:
            status = "× TERUBAH / TIDAK TERDAFTAR"
            baris_dikarantina.append(idx)
            # Save a copy of mismatching uploads into the demo quarantine folder.
            # This does not modify or delete the user's original local file.
            safe_name = f"{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}_{nama_asli}"
            target_path = os.path.join(QUARANTINE_DIR, safe_name)
            uploaded.seek(0)
            with open(target_path, "wb") as out_file:
                while True:
                    chunk = uploaded.read(1024 * 1024)
                    if not chunk:
                        break
                    out_file.write(chunk)
            uploaded.seek(0)
            lokasi_sekarang = target_path

            # The public demo cannot safely infer a reference PDF from an unrelated upload.
            # Compare text only when a matching filename exists in the latest batch memory.
            reference_text = ""
            for item in LATEST_PROCESSED_FILES:
                if item.get("file_name", "").lower() == nama_asli.lower():
                    reference_text = item.get("text", "")
                    break
            if reference_text and teks_baru and not teks_baru.startswith("[Gagal ekstrak"):
                changes = []
                for line in difflib.ndiff(reference_text.splitlines(), teks_baru.splitlines()):
                    if line.startswith("+ "):
                        changes.append(f"[DITAMBAH (+)]: {line[2:]}")
                    elif line.startswith("- "):
                        changes.append(f"[DIHAPUS (-)]: {line[2:]}")
                detail_forensik = " | ".join(changes) if changes else "SHA-256 differs; extracted text appears identical."
            else:
                detail_forensik = "Forensic alert: hash not found in the supplied master list."

        data_detail_laporan.append({
            "no": idx,
            "nama_file": nama_asli,
            "hash": hash_sekarang,
            "status": status,
            "lokasi": lokasi_sekarang,
            "forensik": detail_forensik,
        })

    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    now = datetime.datetime.now()
    hari_indo = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
    bulan_indo = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    waktu = f"{hari_indo[now.weekday()]}, {now.day} {bulan_indo[now.month]} {now.year} | {now.strftime('%H:%M:%S')} WIB"
    summary.append(["AUDIT SYSTEM SUMMARY REPORT", ""])
    summary.append(["DEMO SOFTWARE - FOR DEMONSTRATION ONLY", ""])
    summary.append(["Verification Execution Time:", waktu])
    summary.append(["Total Inspected Files:", len(data_detail_laporan)])
    summary.append(["Total Authentic Files:", total_aman])
    summary.append(["Total Quarantined / Unregistered Files:", len(baris_dikarantina)])
    summary.append(["Quarantined Row Numbers:", str(baris_dikarantina) if baris_dikarantina else "None"])
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 70
    summary["A2"].font = Font(bold=True, color="B91C1C")

    detail = wb.create_sheet("Detail Audit Log")
    detail.append(["No", "File Name", "SHA-256", "Audit Status", "Current Location/Path", "Forensic Text Analysis Result (+/-)"])
    for item in data_detail_laporan:
        detail.append([item["no"], item["nama_file"], item["hash"], item["status"], item["lokasi"], item["forensik"]])
        status_cell = detail.cell(row=detail.max_row, column=4)
        status_cell.fill = PatternFill(fill_type="solid", fgColor="C6EFCE" if "VALID" in item["status"] else "FFC7CE")
    for col, width in {"A": 7, "B": 38, "C": 68, "D": 27, "E": 58, "F": 85}.items():
        detail.column_dimensions[col].width = width
    detail.freeze_panes = "A2"
    pasang_watermark_excel(wb)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), data_detail_laporan, len(baris_dikarantina)


@app.errorhandler(413)
def request_too_large(_error):
    return render_template("index.html", error="Ukuran upload terlalu besar. Coba kurangi jumlah atau ukuran PDF."), 413


@app.route("/", methods=["GET"])
def index():
    return render_template("index.html", data=None, demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS)


@app.route("/proses", methods=["POST"])
def proses_pdf_massal():
    global LATEST_PROCESSED_FILES
    uploaded_files, error = batasi_pdf(request.files.getlist("pdf_files"))
    if error:
        return render_template("index.html", error=error, demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400
    if not uploaded_files:
        return render_template("index.html", error="Pilih minimal satu file PDF yang valid.", demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400

    uploaded_files.sort(key=lambda f: f.filename.lower())
    LATEST_PROCESSED_FILES = []
    display_data = []
    for idx, uploaded in enumerate(uploaded_files, start=1):
        info = proses_guid_dan_ekstrak_teks(uploaded, idx)
        file_hash = hitung_sha256_file(uploaded)
        item = {
            "id": idx,
            "file_name": info["nama_asli"],
            "nama_guid": info["nama_guid"],
            "file_hash": file_hash,
            "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "text": MEMORI_TEKS_ASLI.get(info["guid_key"], ""),
        }
        LATEST_PROCESSED_FILES.append(item)
        display_data.append({"id": idx, "file_name": item["file_name"], "file_hash": file_hash})
    return render_template("index.html", data=display_data, total_files=len(uploaded_files), demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS)


@app.route("/verify", methods=["POST"])
def verifikasi_pdf_vs_csv():
    target_pdfs, error = batasi_pdf(request.files.getlist("verify_pdf"))
    master_file = request.files.get("verify_csv")
    if error:
        return render_template("index.html", v_error=error, demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400
    if not target_pdfs or not master_file or not master_file.filename:
        return render_template("index.html", v_error="Pilih PDF yang akan diverifikasi dan file master XLSX/CSV.", demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400

    try:
        master_name = master_file.filename.lower()
        hash_master_set = set()
        master_file.seek(0)
        if master_name.endswith(".xlsx"):
            master_wb = openpyxl.load_workbook(master_file, read_only=True, data_only=True)
            master_ws = master_wb["DOCUMENT_HASH_REPORT"] if "DOCUMENT_HASH_REPORT" in master_wb.sheetnames else master_wb.active
            for row in master_ws.iter_rows(values_only=True):
                for value in row:
                    if isinstance(value, str) and len(value.strip()) == 64 and all(c in "0123456789abcdefABCDEF" for c in value.strip()):
                        hash_master_set.add(value.strip().lower())
            master_wb.close()
        elif master_name.endswith(".csv"):
            decoded = io.TextIOWrapper(master_file, encoding="utf-8-sig", errors="replace", newline="")
            for row in csv.reader(decoded):
                for value in row:
                    value = (value or "").strip().strip('"')
                    if len(value) == 64 and all(c in "0123456789abcdefABCDEF" for c in value):
                        hash_master_set.add(value.lower())
            master_file.detach() if hasattr(master_file, "detach") else None
        else:
            return render_template("index.html", v_error="File master harus berformat .xlsx atau .csv.", demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400
        if not hash_master_set:
            return render_template("index.html", v_error="Tidak ditemukan hash SHA-256 yang valid di file master.", demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400

        report_bytes, detail_rows, quarantined_count = jalankan_audit_forensik_massal(target_pdfs, hash_master_set)
        view_data = [{
            "file_name": item["nama_file"],
            "file_hash": item["hash"],
            "status": item["status"],
            "is_valid": "VALID" in item["status"],
        } for item in detail_rows]
        response = render_template("index.html", v_data=view_data, ada_palsu=quarantined_count > 0, demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS)
        # Keep report available for explicit download endpoint in this process.

        global LAST_AUDIT_REPORT
        LAST_AUDIT_REPORT = report_bytes
        report_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "demo_software_audit_report.xlsx")
        with open(report_path, "wb") as report_file:
            report_file.write(report_bytes)
        return response

    except Exception as exc:
        app.logger.exception("Verification failed")
        return render_template("index.html", v_error=f"Verifikasi gagal: {exc}", demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 500


LAST_AUDIT_REPORT = None


@app.route("/download-audit-report", methods=["GET"])
def download_audit_report():
    if not LAST_AUDIT_REPORT:
        return render_template("index.html", error="Belum ada laporan verifikasi untuk diunduh.", demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400
    return Response(LAST_AUDIT_REPORT, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=demo_software_audit_report.xlsx"})


@app.route("/download-csv", methods=["GET"])
def download_csv_report():
    if not LATEST_PROCESSED_FILES:
        return render_template("index.html", error="Belum ada batch aktif. Upload PDF terlebih dahulu.", demo_watermark=DEMO_LABEL, max_pdfs=MAX_PDFS), 400
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DOCUMENT_HASH_REPORT"
    ws.append(["Index ID", "File Name", "SHA-256 Cryptographic Hash", "Timestamp"])
    for item in LATEST_PROCESSED_FILES:
        ws.append([item["id"], item["file_name"], item["file_hash"], item["timestamp"]])
    for cell in ws[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for col, width in {"A": 12, "B": 45, "C": 75, "D": 25}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    pasang_watermark_excel(wb)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(output.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=demo_software_master_hash_report.xlsx"})


def buka_browser_otomatis():
    webbrowser.open(f"http://127.0.0.1:{PORT}")


if __name__ == "__main__":
    if os.environ.get("DEMO_AUTO_OPEN", "1") == "1":
        Timer(1, buka_browser_otomatis).start()
    app.run(host="0.0.0.0", debug=False, port=PORT)
